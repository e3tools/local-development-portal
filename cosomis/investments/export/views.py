# Dans MonApp/views.py

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from investments.models import Investment # Assurez-vous d'importer le bon modèle
from django.utils.translation import gettext_lazy as _ # Pour gérer les chaînes traduisibles

def export_investments_to_excel(request):
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    filename = "export_investissements.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    # 2. Création du classeur et de la feuille de calcul
    wb = Workbook()
    ws = wb.active
    ws.title = "Investissements"

    # --- 3. Définition des En-têtes de Colonnes (Headers) ---
    
    columns_lazy = [
        _("Ranking"), _("Title"), _("Description"), _("Responsible Structure"),

        _("ADL Parent"), _("Administrative Level"), _("Sector"), _("Came From (Projects)"),
        
        _("Estimated Cost"), _("Real Cost"), _("Financial Implementation Rate (%)"),
        _("Investment Status"), _("Project Status"),
        
        _("Start Date"), _("Duration (Days)"), _("Delays Consumed (Days)"), 
        _("Physical Execution Rate (%)"),
        
        _("Endorsed by Youth"), _("Endorsed by Women"), _("Endorsed by Agriculturist"), 
        _("Endorsed by Pastoralist"), _("Climate Contribution"), _("Climate Contribution Text"),
        
        _("Latitude"), _("Longitude"), _("Funded By"), _("NoSQL ID"), 
        _("Imported Project ID"),
    ]
    
    columns = [str(header) for header in columns_lazy]

    ws.append(columns)

    investments = Investment.objects.all()

    for investment in investments:

        came_from_list = ", ".join([p.name for p in investment.came_from.all()])
        
        # Récupération du titre du projet financeur (ForeignKey)
        funded_by_title = investment.funded_by.name if investment.funded_by else ""
        
        # Ligne de données correspondante aux colonnes définies à l'étape 3
        row_data = [
            # Descriptif et Administratif
            investment.ranking, investment.title, investment.description, investment.responsible_structure,
            # Relations
            investment.administrative_level.parent.name if investment.administrative_level and investment.administrative_level.parent else "",
            investment.administrative_level.name if investment.administrative_level else "",
            investment.sector.name if investment.sector else "",
            came_from_list,
            # Coûts et Statut
            investment.estimated_cost, investment.real_cost, investment.financial_implementation_rate,
            investment.investment_status,
            investment.project_status,
            # Temps et Exécution
            investment.start_date.strftime("%Y-%m-%d") if investment.start_date else "", # Formatage de la date
            investment.duration, investment.delays_consumed, investment.physical_execution_rate,
            # Impact Social/Climat
            "Oui" if investment.endorsed_by_youth else "Non",
            "Oui" if investment.endorsed_by_women else "Non",
            "Oui" if investment.endorsed_by_agriculturist else "Non",
            "Oui" if investment.endorsed_by_pastoralist else "Non",
            "Oui" if investment.climate_contribution else "Non",
            investment.climate_contribution_text,
            # Géographie et ID
            investment.latitude, investment.longitude, funded_by_title, investment.no_sql_id, 
            investment.imported_project_id,
        ]
        
        ws.append(row_data)

    # Ajustement de la largeur des colonnes (Optionnel) ---
    for col in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 25

    wb.save(response)
    
    return response