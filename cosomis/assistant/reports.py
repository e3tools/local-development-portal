"""Render an assistant answer or report (Markdown) as a Word or PDF file.

Both formats go through the same HTML the chat renders (the model's Markdown,
HTML-escaped first, so nothing it writes is trusted as markup), so what the
user downloads is what they saw in the drawer. Word is built with python-docx
by walking that HTML; PDF is xhtml2pdf over the same HTML plus a small
stylesheet. Relative portal links are made absolute with `base_url` so they
still work once the file leaves the browser.
"""

from __future__ import annotations

import html
import io
from urllib.parse import urljoin

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from django.utils.text import slugify

FORMATS = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    "pdf": ("application/pdf", ".pdf"),
}
DEFAULT_FORMATS = ["docx", "pdf"]
LINK_COLOUR = "00562F"  # --brand-primary-darker


class ReportError(Exception):
    """The file could not be produced (unsupported format or renderer failure)."""


def to_html(text):
    """Same rendering as the chat bubbles: escape first, then Markdown."""
    return markdown.markdown(
        html.escape(text or ""), extensions=["tables", "nl2br", "sane_lists"])


def filename(title, fmt):
    return (slugify(title)[:60] or "rapport") + FORMATS[fmt][1]


def content_type(fmt):
    return FORMATS[fmt][0]


def render(title, body_markdown, fmt, base_url="", meta="", footer=""):
    """Return the bytes of `body_markdown` as `fmt`, with `title` as heading.

    `meta` (who/when) goes under the title, `footer` (the disclaimer) at the
    end. `base_url` absolutises the relative portal links the tools return.
    """
    if fmt not in FORMATS:
        raise ReportError(f"Unsupported format '{fmt}'.")
    soup = BeautifulSoup(to_html(body_markdown), "html.parser")
    if base_url:
        for anchor in soup.find_all("a", href=True):
            anchor["href"] = urljoin(base_url, anchor["href"])
    if fmt == "docx":
        return _docx(title, soup, meta, footer)
    return _pdf(title, soup, meta, footer)


# -- Word -------------------------------------------------------------------

def _docx(title, soup, meta, footer):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    doc.add_heading(title, level=0)
    if meta:
        run = doc.add_paragraph().add_run(meta)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    for node in soup.contents:
        _docx_block(doc, node, list_level=0)
    if footer:
        run = doc.add_paragraph().add_run(footer)
        run.italic = True
        run.font.size = Pt(9)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _docx_block(doc, node, list_level):
    if isinstance(node, NavigableString):
        text = str(node).strip()
        if text:
            doc.add_paragraph(text)
        return
    if not isinstance(node, Tag):
        return
    name = node.name
    if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
        # Level 0 is the document title, so the body's h1 becomes level 1.
        doc.add_heading(node.get_text(" ", strip=True), level=min(int(name[1]), 4))
    elif name == "p":
        _docx_inline(doc.add_paragraph(), node)
    elif name in ("ul", "ol"):
        style = "List Bullet" if name == "ul" else "List Number"
        if list_level:
            style += f" {min(list_level + 1, 3)}"
        for item in node.find_all("li", recursive=False):
            paragraph = doc.add_paragraph(style=style)
            nested = []
            for child in item.children:
                if isinstance(child, Tag) and child.name in ("ul", "ol"):
                    nested.append(child)
                else:
                    _docx_inline(paragraph, child)
            for sub in nested:
                _docx_block(doc, sub, list_level + 1)
    elif name == "table":
        rows = node.find_all("tr")
        cols = max((len(r.find_all(["th", "td"])) for r in rows), default=0)
        if not rows or not cols:
            return
        table = doc.add_table(rows=0, cols=cols)
        table.style = "Table Grid"
        for row in rows:
            cells = table.add_row().cells
            for i, cell in enumerate(row.find_all(["th", "td"])[:cols]):
                paragraph = cells[i].paragraphs[0]
                _docx_inline(paragraph, cell, bold=(cell.name == "th"))
    elif name == "blockquote":
        _docx_inline(doc.add_paragraph(style="Quote"), node)
    elif name == "pre":
        _docx_inline(doc.add_paragraph(), node, code=True)
    elif name == "hr":
        doc.add_paragraph()
    else:
        for child in node.children:
            _docx_block(doc, child, list_level)


def _docx_inline(paragraph, node, bold=False, italic=False, code=False):
    if isinstance(node, NavigableString):
        text = str(node)
        if text.strip() or " " in text:
            run = paragraph.add_run(text.replace("\n", " "))
            run.bold = bold or None
            run.italic = italic or None
            if code:
                run.font.name = "Consolas"
        return
    if not isinstance(node, Tag):
        return
    name = node.name
    if name == "br":
        paragraph.add_run().add_break()
    elif name == "a" and node.get("href"):
        _docx_hyperlink(paragraph, node.get_text(" ", strip=True) or node["href"], node["href"])
    else:
        bold = bold or name in ("strong", "b")
        italic = italic or name in ("em", "i")
        code = code or name in ("code", "pre")
        for child in node.children:
            _docx_inline(paragraph, child, bold=bold, italic=italic, code=code)


def _docx_hyperlink(paragraph, text, url):
    """python-docx has no hyperlink API; build the run in raw OOXML."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), LINK_COLOUR)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(colour)
    props.append(underline)
    run.append(props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)  # noqa: SLF001 — the documented way to append raw XML


# -- PDF --------------------------------------------------------------------

PDF_CSS = """
@page { size: A4; margin: 2cm 1.8cm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; color: #111827; line-height: 1.4; }
h1.report-title { font-size: 20pt; color: #00562f; margin: 0 0 4pt; }
p.report-meta { color: #6b7280; font-size: 9pt; margin: 0 0 14pt; }
h1 { font-size: 15pt; color: #00562f; margin: 14pt 0 6pt; }
h2 { font-size: 13pt; color: #00562f; margin: 12pt 0 5pt; }
h3, h4 { font-size: 11pt; color: #00562f; margin: 10pt 0 4pt; }
p { margin: 0 0 6pt; }
ul, ol { margin: 0 0 6pt 14pt; }
table { width: 100%; border-collapse: collapse; margin: 4pt 0 8pt; font-size: 9.5pt; }
th, td { border: 1px solid #cbd5d1; padding: 3pt 5pt; text-align: left; vertical-align: top; }
th { background-color: #eaf4ec; color: #00562f; }
a { color: #00562f; text-decoration: underline; }
code { font-family: Courier, monospace; font-size: 9pt; }
p.report-footer { color: #6b7280; font-size: 8.5pt; margin-top: 16pt; }
"""


def _pdf(title, soup, meta, footer):
    from xhtml2pdf import pisa

    document = (
        "<html><head><meta charset='utf-8'><style>" + PDF_CSS + "</style></head><body>"
        f"<h1 class='report-title'>{html.escape(title)}</h1>"
        + (f"<p class='report-meta'>{html.escape(meta)}</p>" if meta else "")
        + str(soup)
        + (f"<p class='report-footer'>{html.escape(footer)}</p>" if footer else "")
        + "</body></html>"
    )
    out = io.BytesIO()
    status = pisa.CreatePDF(io.StringIO(document), dest=out, encoding="utf-8")
    if status.err:
        raise ReportError("The PDF renderer reported an error.")
    return out.getvalue()
